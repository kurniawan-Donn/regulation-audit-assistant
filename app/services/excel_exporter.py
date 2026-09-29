"""
Excel Exporter Service (Fase 12).

10 kolom sesuai spesifikasi kertas kerja audit:
No | Referensi | Aspek | BAB/Pasal/Ayat | Ketentuan | Prosedur Pemeriksaan |
Metode Audit | Kebutuhan Bukti | Status Kepatuhan | Catatan/Temuan

Kolom "Status Kepatuhan Bank" diberi dropdown data validation.
Kolom auditor (Status + Catatan) diberi warna latar khusus auditor.
Baris flagged (is_valid=False) tetap diblok warna merah muda.
"""

import re
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.worksheet import Worksheet

from app.prompts.aspects import STATUS_KEPATUHAN

# ---- Styling ----
HEADER_FILL = PatternFill(start_color="1E2761", end_color="1E2761", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)
FLAGGED_FILL = PatternFill(start_color="FBEAE8", end_color="FBEAE8", fill_type="solid")
AUDITOR_FILL = PatternFill(start_color="FFF8E1", end_color="FFF8E1", fill_type="solid")

COLUMNS: list[tuple[str, int]] = [
    ("No", 5),
    ("Referensi Regulasi", 22),
    ("Aspek / Domain", 22),
    ("BAB / Pasal / Ayat", 20),
    ("Ketentuan / Poin Regulasi", 42),
    ("Prosedur Pemeriksaan", 48),
    ("Metode Audit", 20),
    ("Kebutuhan Bukti", 34),
    ("Status Kepatuhan Bank", 18),
    ("Catatan / Temuan", 32),
]

# Index (1-based) kolom auditor — untuk styling + data validation
_COL_STATUS = 9
_COL_CATATAN = 10


def _safe_filename_component(text: str) -> str:
    text = (text or "").strip() or "regulasi"
    text = re.sub(r"[^\w\-.]+", "_", text)
    return text.strip("_")[:80] or "regulasi"


def build_export_filename(referensi_regulasi: str) -> str:
    return f"audit_checklist_{_safe_filename_component(referensi_regulasi)}.xlsx"


def export_results_to_excel(results: list[dict[str, Any]], output_path: str) -> str:
    wb = Workbook()
    ws: Worksheet = wb.active
    ws.title = "Audit Checklist"

    # --- Header ---
    for col_idx, (header, width) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(col_idx)].width = width
    ws.row_dimensions[1].height = 26

    # --- Baris data ---
    for i, item in enumerate(results, start=1):
        is_valid = item.get("is_valid", True)
        row_idx = i + 1

        row_values = [
            i,                                       # 1
            item.get("referensi_regulasi", ""),      # 2
            item.get("aspek", ""),                   # 3
            item.get("bab_pasal_ayat", ""),          # 4
            item.get("point_ketentuan", ""),         # 5
            item.get("pemeriksaan", ""),             # 6
            item.get("metode_audit", ""),            # 7
            item.get("kebutuhan_bukti", ""),         # 8
            item.get("status_kepatuhan", "Belum Diaudit"),  # 9
            item.get("catatan_temuan", ""),          # 10
        ]

        for col_idx, value in enumerate(row_values, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            if col_idx in (1, _COL_STATUS):
                cell.alignment = Alignment(horizontal="center", vertical="top",
                                           wrap_text=(col_idx == _COL_STATUS))
            else:
                cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)

            # Prioritas warna: flagged > auditor > netral
            if not is_valid:
                cell.fill = FLAGGED_FILL
            elif col_idx in (_COL_STATUS, _COL_CATATAN):
                cell.fill = AUDITOR_FILL

    # --- Data validation dropdown untuk kolom Status Kepatuhan ---
    options = ",".join(STATUS_KEPATUHAN)
    dv = DataValidation(
        type="list",
        formula1=f'"{options}"',
        allow_blank=True,
        showDropDown=False,  # False = dropdown tetap muncul di Excel
    )
    dv.error = "Pilih salah satu status yang tersedia."
    dv.errorTitle = "Status tidak valid"
    ws.add_data_validation(dv)
    last_row_for_validation = max(2, len(results) + 1)
    dv.add(f"{get_column_letter(_COL_STATUS)}2:{get_column_letter(_COL_STATUS)}{last_row_for_validation}")

    # --- Freeze & auto filter ---
    ws.freeze_panes = "A2"
    last_col_letter = get_column_letter(len(COLUMNS))
    last_row = len(results) + 1
    ws.auto_filter.ref = f"A1:{last_col_letter}{last_row}"

    Path(output_path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    return output_path