"""
Excel Exporter untuk Gap Analysis Tool.

Output: 1 sheet "Persyaratan" dengan kolom:
No | Persyaratan | Referensi | Catatan | Status Gap

Status Gap dikosongkan (auditor isi manual di Excel atau web).
"""

import re
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.worksheet import Worksheet

HEADER_FILL = PatternFill(start_color="1E2761", end_color="1E2761", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)
STATUS_FILL = PatternFill(start_color="FFF8E1", end_color="FFF8E1", fill_type="solid")

COLUMNS: list[tuple[str, int]] = [
    ("No", 5),
    ("Persyaratan", 45),
    ("Referensi", 40),
    ("Catatan", 80),
    ("Status Gap", 18),
]

STATUS_GAP_OPTIONS = ["Sesuai", "Sebagian Sesuai", "Tidak Sesuai", "Belum Dinilai", "N/A"]
_COL_STATUS = 5


def _safe_filename_component(text: str) -> str:
    text = (text or "").strip() or "regulasi"
    text = re.sub(r"[^\w\-.]+", "_", text)
    return text.strip("_")[:80] or "regulasi"


def build_requirements_filename(referensi_regulasi: str) -> str:
    return f"gap_analysis_{_safe_filename_component(referensi_regulasi)}.xlsx"


def _strip_markdown(text: str) -> str:
    """
    Buang markdown bold (**text**) untuk kolom Excel fase ini.
    Multi-line tetap dipertahankan (\n).
    """
    if not text:
        return ""
    # Buang **bold**
    text = re.sub(r"\*\*(.+?)\*\*", r"\1", text)
    return text.strip()


def export_requirements_to_excel(
    requirements: list[dict[str, Any]],
    output_path: str,
) -> str:
    """Tulis daftar persyaratan ke file xlsx."""
    wb = Workbook()
    ws: Worksheet = wb.active
    ws.title = "Persyaratan"

    # Header
    for col_idx, (header, width) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(col_idx)].width = width
    ws.row_dimensions[1].height = 24

    # Baris data
    for i, item in enumerate(requirements, start=1):
        row_idx = i + 1
        values = [
            i,
            item.get("persyaratan", ""),
            item.get("referensi", ""),
            _strip_markdown(item.get("catatan", "")),
            item.get("status_gap", "") or "",
        ]

        for col_idx, value in enumerate(values, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            if col_idx == 1:
                cell.alignment = Alignment(horizontal="center", vertical="top")
            elif col_idx == _COL_STATUS:
                cell.alignment = Alignment(horizontal="center", vertical="top", wrap_text=True)
                cell.fill = STATUS_FILL
            else:
                cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)

    # Data validation dropdown untuk kolom Status
    options = ",".join(STATUS_GAP_OPTIONS)
    dv = DataValidation(
        type="list",
        formula1=f'"{options}"',
        allow_blank=True,
        showDropDown=False,
    )
    dv.error = "Pilih salah satu status."
    dv.errorTitle = "Status tidak valid"
    ws.add_data_validation(dv)
    last_row = max(2, len(requirements) + 1)
    dv.add(f"{get_column_letter(_COL_STATUS)}2:{get_column_letter(_COL_STATUS)}{last_row}")

    # Freeze & filter
    ws.freeze_panes = "A2"
    last_col_letter = get_column_letter(len(COLUMNS))
    ws.auto_filter.ref = f"A1:{last_col_letter}{last_row}"

    Path(output_path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    return output_path