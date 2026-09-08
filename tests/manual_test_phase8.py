import io
import os
import sys
from unittest.mock import patch

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from openpyxl import load_workbook

from app.services.excel_exporter import build_export_filename, export_results_to_excel

SAMPLE_TXT_PATH = os.path.join(os.path.dirname(__file__), "sample_files", "sample_regulation.txt")


def print_section(title):
    print("\n" + "=" * 60)
    print(title)
    print("=" * 60)


SAMPLE_RESULTS = [
    {
        "referensi_regulasi": "PERATURAN CONTOH 01/2026",
        "bab_pasal_ayat": "BAB V - Pasal 20 - (1)",
        "point_ketentuan": "Kewajiban kebijakan keamanan informasi",
        "pemeriksaan": "Periksa keberadaan kebijakan keamanan informasi yang disahkan.",
        "is_valid": True,
        "validation_notes": [],
    },
    {
        "referensi_regulasi": "PERATURAN CONTOH 01/2026",
        "bab_pasal_ayat": "Pasal 999",
        "point_ketentuan": "Ketentuan yang dikarang AI",
        "pemeriksaan": "Periksa sesuatu yang tidak ada di sumber.",
        "is_valid": False,
        "validation_notes": ["Referensi BAB/Pasal/Ayat tidak ditemukan persis di teks sumber - perlu ditinjau manual"],
    },
]


def test_build_export_filename():
    print_section("TEST 1: build_export_filename()")
    name = build_export_filename("POJK 11/POJK.03/2022")
    print(f"Hasil: {name}")
    assert name.startswith("audit_checklist_")
    assert name.endswith(".xlsx")
    assert "/" not in name, "Nama file tidak boleh mengandung karakter path separator"
    print("OK - nama file aman dipakai di semua OS")


def test_export_and_reopen_excel():
    print_section("TEST 2: export_results_to_excel() lalu dibuka ulang & diverifikasi")
    output_path = "/tmp/test_audit_checklist.xlsx"
    export_results_to_excel(SAMPLE_RESULTS, output_path)

    assert os.path.exists(output_path)
    wb = load_workbook(output_path)
    assert "Audit Checklist" in wb.sheetnames
    ws = wb["Audit Checklist"]

    # Header
    header_row = [cell.value for cell in ws[1]]
    print(f"Header: {header_row}")
    assert header_row == ["No", "Referensi Regulasi", "BAB/Pasal/Ayat", "Point/Ketentuan", "Pemeriksaan", "Status", "Catatan Validasi"]

    # Baris data
    assert ws.cell(row=2, column=1).value == 1
    assert ws.cell(row=2, column=6).value == "Valid"
    assert ws.cell(row=3, column=1).value == 2
    assert ws.cell(row=3, column=6).value == "Perlu Ditinjau"
    assert "tidak ditemukan" in ws.cell(row=3, column=7).value
    print("OK - isi baris data & status Valid/Perlu Ditinjau sesuai")

    # Freeze panes & auto filter
    assert ws.freeze_panes == "A2"
    assert ws.auto_filter.ref == f"A1:G{len(SAMPLE_RESULTS) + 1}"
    print(f"OK - freeze_panes={ws.freeze_panes}, auto_filter={ws.auto_filter.ref}")

    # Header bold & fill
    assert ws.cell(row=1, column=1).font.bold is True
    print("OK - header bold")

    # Baris flagged punya fill berbeda dari baris valid
    fill_valid = ws.cell(row=2, column=3).fill.start_color.rgb
    fill_flagged = ws.cell(row=3, column=3).fill.start_color.rgb
    assert fill_valid != fill_flagged
    print(f"OK - baris flagged ({fill_flagged}) punya highlight berbeda dari baris valid ({fill_valid})")

    os.remove(output_path)


@patch("app.api.routes.analyze_chunks_batch")
def test_export_endpoint_end_to_end(mock_analyze):
    print_section("TEST 3: endpoint /api/analyze -> /api/export end-to-end")
    from fastapi.testclient import TestClient
    from app.main import app

    mock_analyze.return_value = [
        {
            "referensi_regulasi": "PERATURAN CONTOH 01/2026",
            "bab_pasal_ayat": "BAB V - Pasal 20 - (1)",
            "point_ketentuan": "Kewajiban kebijakan keamanan informasi",
            "pemeriksaan": "Periksa keberadaan kebijakan keamanan informasi yang disahkan.",
        }
    ]

    client = TestClient(app)
    with open(SAMPLE_TXT_PATH, "rb") as f:
        analyze_res = client.post(
            "/api/analyze",
            files={"file": ("sample_regulation.txt", f, "text/plain")},
            data={"nama_regulasi": "PERATURAN CONTOH 01/2026"},
        )
    assert analyze_res.status_code == 200, analyze_res.text
    result_id = analyze_res.json()["result_id"]
    print(f"result_id: {result_id}")

    export_res = client.get(f"/api/export/{result_id}")
    assert export_res.status_code == 200, export_res.text
    assert export_res.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert "attachment" in export_res.headers.get("content-disposition", "")
    print(f"Content-Disposition: {export_res.headers.get('content-disposition')}")

    # Buka file Excel dari response bytes untuk verifikasi isinya benar2 valid
    wb = load_workbook(io.BytesIO(export_res.content))
    ws = wb["Audit Checklist"]
    assert ws.cell(row=1, column=1).value == "No"
    assert ws.cell(row=2, column=2).value == "PERATURAN CONTOH 01/2026"
    print("OK - file Excel hasil endpoint valid dan bisa dibuka ulang dengan openpyxl")


def test_export_not_found():
    print_section("TEST 4: export result_id yang tidak ada -> 404")
    from fastapi.testclient import TestClient
    from app.main import app

    client = TestClient(app)
    res = client.get("/api/export/id-tidak-ada")
    assert res.status_code == 404
    print("OK - 404 untuk result_id yang tidak ditemukan")


if __name__ == "__main__":
    test_build_export_filename()
    test_export_and_reopen_excel()
    test_export_endpoint_end_to_end()
    test_export_not_found()
    print_section("SEMUA TEST PHASE 8 SELESAI")
